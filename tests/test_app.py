"""Tests for RubriPDF: PDF generation, grade extraction with its quality
checks, CSV-injection hardening, and spreadsheet/CSV parsing."""

import io
import os
import sys

import pytest
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import app  # noqa: E402


# ── helpers ──────────────────────────────────────────────────────────────────

def make_config(**over):
    cfg = {
        'title': 'Práctica 1', 'subject': 'Materia', 'year': '2025/26',
        'language': 'galician', 'color': '210,235,230',
        'obs_pages': 0, 'obs_template': False,
        'criteria': [{'description': 'c1', 'points': 2},
                     {'description': 'c2', 'points': 5}],
        'students': ['Noa Fernández', 'María López'],
    }
    cfg.update(over)
    return cfg


def gen(**over):
    return app.generate_pdf_python(make_config(**over))


def fields_of(pdf_bytes):
    return PdfReader(io.BytesIO(pdf_bytes)).get_fields() or {}


def fill(pdf_bytes, text=None, checks=None):
    """Return a BytesIO of the PDF with text fields set and checkboxes ticked."""
    reader = PdfReader(io.BytesIO(pdf_bytes))
    writer = PdfWriter(clone_from=reader)
    if text:
        for pg in writer.pages:
            writer.update_page_form_field_values(pg, text, auto_regenerate=False)
    if checks:
        for pg in writer.pages:
            for a in (pg.get('/Annots') or []):
                o = a.get_object()
                if str(o.get('/T')) in checks:
                    o[NameObject('/V')] = NameObject('/Yes')
                    o[NameObject('/AS')] = NameObject('/Yes')
    buf = io.BytesIO()
    writer.write(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def client():
    app.app.config['TESTING'] = True
    return app.app.test_client()


def do_extract(client, buf, detailed=False, lang=None):
    data = {'file': (buf, 'e.pdf')}
    if detailed:
        data['detailed'] = '1'
    if lang:
        data['lang'] = lang
    return client.post('/extract', data=data).get_json()


def warn_types(res):
    return [w['t'] for w in res['warnings']]


# ── PDF generation ───────────────────────────────────────────────────────────

def test_page_count():
    r = PdfReader(io.BytesIO(gen(obs_pages=1, students=['A', 'B', 'C'])))
    assert len(r.pages) == 1 + 3 * (1 + 1)  # cover + 3 * (student page + 1 obs)


def test_hidden_name_and_cmax_and_np_fields():
    f = fields_of(gen())
    assert str(f['name1'].get('/V')) == 'Noa Fernández'
    assert str(f['name2'].get('/V')) == 'María López'
    assert {str(f['cmax0'].get('/V')), str(f['cmax1'].get('/V'))} == {'2', '5'}
    assert 'np1' in f and 'np2' in f


def test_unicode_names_do_not_crash():
    pdf = gen(students=['Zażółć gęślą jaźń', 'Иван Петров', 'Núria'])
    assert len(pdf) > 1000
    f = fields_of(pdf)
    assert str(f['name2'].get('/V')) == 'Иван Петров'


def test_acroform_dr_and_embedded_font():
    raw = gen()
    assert b'/DR' in raw and b'/Helv' in raw and b'/ZaDb' in raw
    assert b'DejaVu' in raw
    assert b'/NeedAppearances true' in raw


def test_obs_template_toggle():
    on = fields_of(gen(obs_template=True))['obs1p0'].get('/V')
    assert '### 1. c1' in str(on)
    off = fields_of(gen(obs_template=False))['obs1p0'].get('/V')
    assert off in (None, '')


def test_obs_pages_zero_keeps_first_box_only():
    r = PdfReader(io.BytesIO(gen(obs_pages=0, students=['A'])))
    assert len(r.pages) == 2  # cover + single student page
    f = fields_of(gen(obs_pages=0, students=['A']))
    assert 'obs1p0' in f and 'obs1p1' not in f


# ── extraction: quality checks ───────────────────────────────────────────────

def test_extract_normal(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5', 'final1': '7'})
    res = do_extract(client, buf)
    assert res['csv'].splitlines()[1] == 'Noa,7'
    assert res['warnings'] == []


def test_extract_autofill_blank_final(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5'})  # final blank
    res = do_extract(client, buf)
    assert res['csv'].splitlines()[1] == 'Noa,7'
    assert warn_types(res) == ['final_autofill']


def test_extract_decimal_and_delimiter_galician(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2,5', 'pts1c1': '5', 'final1': '7,5'})
    res = do_extract(client, buf, lang='gl')
    assert res['csv'].splitlines()[1] == 'Noa;7,5'       # "," decimal, ";" delimiter


def test_extract_decimal_and_delimiter_english(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2.5', 'pts1c1': '5', 'final1': '7.5'})
    res = do_extract(client, buf, lang='en')
    assert res['csv'].splitlines()[1] == 'Noa,7.5'       # "." decimal, "," delimiter


def test_extract_typed_decimal_normalised_to_locale(client):
    # a "." typed by the teacher is re-formatted to the language's decimal
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2.5', 'pts1c1': '5', 'final1': '7.5'})
    res = do_extract(client, buf, lang='gl')
    assert res['csv'].splitlines()[1] == 'Noa;7,5'


def test_extract_autofill_uses_locale_decimal(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2,5', 'pts1c1': '5'})     # final blank -> sum 7.5
    res = do_extract(client, buf, lang='gl')
    assert res['csv'].splitlines()[1] == 'Noa;7,5'
    assert 'final_autofill' in warn_types(res)


def test_extract_mismatch(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5', 'final1': '6'})
    res = do_extract(client, buf)
    assert 'final_mismatch' in warn_types(res)


def test_extract_empty_criterion(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'final1': '2'})  # c1 (index 1) blank
    res = do_extract(client, buf)
    ws = [w for w in res['warnings'] if w['t'] == 'crit_empty']
    assert ws and ws[0]['c'] == [1]


def test_extract_final_non_numeric(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5', 'final1': 'aprobado'})
    res = do_extract(client, buf)
    assert 'final_nan' in warn_types(res)


def test_extract_over_max(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '9', 'pts1c1': '5', 'final1': '14'})  # c0 max is 2
    res = do_extract(client, buf)
    over = [w for w in res['warnings'] if w['t'] == 'crit_over']
    assert over and over[0]['c'] == 0 and over[0]['max'] == 2.0


def test_extract_np_empty(client):
    pdf = gen(students=['Noa', 'María'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5', 'final1': '7'}, checks={'np2'})
    res = do_extract(client, buf)
    lines = res['csv'].splitlines()
    assert lines[1] == 'Noa,7' and lines[2] == 'María,NP'
    assert res['warnings'] == []


def test_extract_np_with_grades_warns(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2'}, checks={'np1'})
    res = do_extract(client, buf)
    assert res['csv'].splitlines()[1] == 'Noa,NP'
    assert warn_types(res) == ['np_has_grades']


def test_extract_detailed_columns(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5', 'final1': '7'})
    res = do_extract(client, buf, detailed=True)
    header = res['csv'].splitlines()[0]
    assert header == 'Student,Final grade,C1 Level,C1 pts,C2 Level,C2 pts'
    assert res['csv'].splitlines()[1].endswith(',2,,5')


def test_extract_header_localised(client):
    pdf = gen(students=['Noa'])
    buf = fill(pdf, {'pts1c0': '2', 'pts1c1': '5', 'final1': '7'})
    res = do_extract(client, buf, detailed=True, lang='gl')
    assert res['csv'].splitlines()[0] == 'Alumno/a;Nota final;C1 Nivel;C1 ptos;C2 Nivel;C2 ptos'


def test_extract_derives_points_from_level(client):
    # derive_points on: ticking level 4 (fraction 1.0) on a 5-pt criterion -> 5
    pdf = gen(students=['Noa'], derive_points=True,
              criteria=[{'description': 'c1', 'points': 5}])
    buf = fill(pdf, checks={'lvl1c0_4'})          # student 1, criterion 0, top level (Excelente)
    res = do_extract(client, buf, detailed=True, lang='gl')
    lines = res['csv'].splitlines()
    assert lines[0] == 'Alumno/a;Nota final;C1 Nivel;C1 ptos'
    assert lines[1] == 'Noa;5;5;5'                # level 5 shown, 5 pts derived (1.0×5), final = sum

    res_es = do_extract(client, fill(gen(students=['Noa']), {'final1': '7'}), lang='es')
    assert res_es['csv'].splitlines()[0] == 'Alumno/a;Nota final'


def test_extract_csv_injection_sanitised(client):
    pdf = gen(students=['=HYPERLINK("x")', '@SUM(A1)', 'María López'])
    buf = fill(pdf, {'final1': '5', 'final2': '6', 'final3': '7'})
    res = do_extract(client, buf)
    body = res['csv'].splitlines()[1:]
    assert body[0].startswith('"\'=HYPERLINK') or body[0].startswith("'=HYPERLINK")
    assert body[1].startswith("'@SUM(A1)")
    assert body[2].startswith('María López')  # untouched


def test_extract_no_fields_errors(client):
    from fpdf import FPDF
    p = FPDF()
    p.add_page()
    p.set_font('Helvetica', size=12)
    p.cell(0, 10, 'no form here')
    blank = io.BytesIO(bytes(p.output()))
    r = client.post('/extract', data={'file': (blank, 'x.pdf')})
    assert r.status_code == 400 and r.get_json()['error'] == 'no_fields'


# ── parsing ──────────────────────────────────────────────────────────────────

def _parse_csv(client, text, filename='a.csv', encoding='utf-8'):
    buf = io.BytesIO(text.encode(encoding))
    return client.post('/parse-csv', data={'file': (buf, filename)}).get_json()


def test_parse_csv_basic(client):
    d = _parse_csv(client, 'nombre\nAna\nBea\n')     # header label skipped
    assert d['students'] == ['Ana', 'Bea']
    assert d['format'] == 'generic'


def test_help_route(client):
    for lang in ('gl', 'es', 'en', 'pt'):
        r = client.get('/help/' + lang)
        assert r.status_code == 200
        assert '<h5' in r.get_data(as_text=True)
    assert client.get('/help/zz').status_code == 200     # unknown -> English


def test_md_to_html():
    h = app._md_to_html('## T\ntexto **b** e *i*.\n\n- a\n- b\n\n![x](/static/y.png)')
    assert '<h5' in h and '<strong>b</strong>' in h and '<em>i</em>' in h
    assert '<ul><li>a</li><li>b</li></ul>' in h
    assert '<img src="/static/y.png"' in h


def test_parse_csv_one_name_per_line(client):
    # single column: names keep their internal comma, none split or dropped
    d = _parse_csv(client, 'Seoane, Icía\nGondar Veiga, Brais\nMéndez Añón, Iago\n', 'lista.txt')
    assert d['count'] == 3
    assert d['students'] == ['Seoane, Icía', 'Gondar Veiga, Brais', 'Méndez Añón, Iago']


def test_parse_csv_headerless_keeps_first(client):
    # the old DictReader swallowed the first line as a header; it must not now
    d = _parse_csv(client, 'Rega Souto, Antón\nTorres da Ponte, Uxía\n')
    assert d['students'] == ['Rega Souto, Antón', 'Torres da Ponte, Uxía']


def test_parse_csv_windows_crlf(client):
    # a .txt/.csv saved on Windows carries \r\n — the \r must not survive
    d = _parse_csv(client, 'Seoane, Icía\r\nGondar Veiga, Brais\r\n', 'win.txt')
    assert d['students'] == ['Seoane, Icía', 'Gondar Veiga, Brais']


def test_parse_csv_skips_header_blank_quotes_ordinal(client):
    # BOM + header label + blank line + fully-quoted value + bare ordinal
    d = _parse_csv(client, 'Nome\r\n\r\n"Seoane, Icía"\r\nGondar Veiga, Brais\r\nNº\r\n',
                   'l.csv', encoding='utf-8-sig')
    assert d['students'] == ['Seoane, Icía', 'Gondar Veiga, Brais']


def test_parse_xls_preserves_no_nu_names(client):
    # HTML-disguised .xls with the name in the first column, incl. No*/Nu* names
    html = (
        '<html><body><table>'
        '<tr><td>Nome</td><td>x</td></tr>'
        '<tr><td>Noa Fernández</td><td></td></tr>'
        '<tr><td>Nuria Castro</td><td></td></tr>'
        '<tr><td>Núria Pons</td><td></td></tr>'
        '<tr><td>Nº</td><td></td></tr>'          # bare ordinal, must be dropped
        '</table></body></html>'
    ).encode('utf-8')
    d = client.post('/parse-xls', data={'file': (io.BytesIO(html), 'l.xls')}).get_json()
    assert d['students'] == ['Noa Fernández', 'Nuria Castro', 'Núria Pons']


# Invented data mimicking the structure of the USC "Secretaría Virtual"
# student-list export (HTML-disguised .xls): an outer wrapper table nesting the
# real one, a title row, the real header at row 3, an empty "Modo" column, and a
# "Nº de alumnos/as" footer. Names carry the source's own quirks — "APELIDOS,
# Nome" order, ALL-CAPS, accents/ñ, a stray space before the comma, and names
# starting with No*/Nu* — which MUST survive verbatim (this list is matched
# against other sources, so no normalisation is allowed).
USC_STUDENTS = [
    'Seoane , Icía',                # stray space before the comma
    'GONDAR VEIGA, BRAIS',          # ALL-CAPS
    'Méndez Añón, Iago',            # accents and ñ
    'Nogareda Mar, Noa',            # starts with "No" — not a bare ordinal
    'Nunes Vila, Xoel',             # starts with "Nu"
    'Rega Souto, Antón',
    'Torres da Ponte, Ana María',   # particle + two given names
]


def _usc_xls_bytes(students):
    body = ''.join(
        f'<tr><td>{i:08d}X</td><td>{n}</td><td>FB</td><td></td>'
        f'<td>Módulo de Formación Básica</td></tr>'
        for i, n in enumerate(students, 1)
    )
    inner = (
        '<tr><td>Listaxe de alumnos e alumnas</td></tr>'
        '<tr><td>Documento</td><td>Nome</td><td>Carácter</td>'
        '<td>Modo</td><td>Módulo</td></tr>'
        + body +
        f'<tr><td>Nº de alumnos/as : {len(students)}</td></tr>'
    )
    # Outer wrapper table nesting the real one, as the USC export does.
    return (
        '<html><head>'
        '<meta http-equiv="Content-Type" content="text/html; charset=utf-8">'
        '</head><body><table><tr><td>'
        f'<table>{inner}</table>'
        '</td></tr></table></body></html>'
    ).encode('utf-8')


def test_parse_usc_virtual_secretary_list(client):
    xls = _usc_xls_bytes(USC_STUDENTS)
    d = client.post('/parse-xls',
                    data={'file': (io.BytesIO(xls), 'Listaxe de alumnos.xls')}).get_json()
    assert d['selected_column'] == 'Nome'
    assert d['format'] == 'usc'
    assert d['columns'] == ['Documento', 'Nome', 'Carácter', 'Modo', 'Módulo']
    # Every student read, none lost to the nested-table/title/footer noise, and
    # each name byte-for-byte identical to the source (no normalisation).
    assert d['students'] == USC_STUDENTS


def test_parse_moodle_split_name(client):
    # Moodle "participants" export: name split across Nome/Apelidos, with an
    # email column and a trailing empty Grupos. Must auto-detect and assemble
    # "Apelidos, Nome" — the same shape as the USC list — with no manual step.
    import openpyxl
    from io import BytesIO
    wb = openpyxl.Workbook()
    ws = wb.active
    # Invented data — never copy names/emails from a real export.
    ws.append(['Nome', 'Apelidos', 'Enderezo de correo', 'Grupos'])
    ws.append(['Icía', 'Seoane', 'icia.seoane@example.com'])
    ws.append(['Brais', 'Gondar Veiga', 'brais.gondar@example.com'])
    ws.append(['Noa', 'Nogareda Mar', 'noa.nogareda@example.com'])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    d = client.post('/parse-xls',
                    data={'file': (buf, 'courseid_1_participants.xlsx')}).get_json()
    assert d['format'] == 'moodle'
    assert d['count'] == 3
    assert d['selected_column'] == 'Apelidos, Nome'
    assert 'note' not in d                       # labels recognised -> no warning
    assert d['students'] == ['Seoane, Icía', 'Gondar Veiga, Brais', 'Nogareda Mar, Noa']


def test_parse_csv_reports_duplicates(client):
    d = _parse_csv(client, 'Seoane, Icía\nGondar Veiga, Brais\nSeoane, Icía\n', 'l.txt')
    assert d['count'] == 3
    assert d['duplicates'] == 1                   # one repeated entry, none removed
    assert d['students'].count('Seoane, Icía') == 2


def test_parse_csv_no_duplicates(client):
    d = _parse_csv(client, 'Seoane, Icía\nGondar Veiga, Brais\n', 'l.txt')
    assert d['duplicates'] == 0


def test_parse_csv_flags_multicolumn(client):
    # a Moodle-style multi-column CSV uploaded by mistake: imported whole, flagged
    d = _parse_csv(client, 'Nome;Apelidos;Correo\nIcía;Seoane;a@example.com\n'
                           'Brais;Gondar;b@example.com\n', 'moodle.csv')
    assert d['note'] == 'looks_multicolumn'


def test_parse_csv_comma_names_not_flagged(client):
    # a legit "Apelidos, Nome" list must NOT trigger the multi-column warning
    d = _parse_csv(client, 'Seoane, Icía\nGondar Veiga, Brais\nMéndez Añón, Iago\n')
    assert 'note' not in d


def test_parse_moodle_unknown_language_headers(client):
    # Headers in a language we do NOT list. The email column anchors detection
    # (language-independent) and Moodle's fixed given/surname/email order lets us
    # assemble the name — but flagged 'inferred' so the UI asks the user to check.
    import openpyxl
    from io import BytesIO
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Izena', 'Abizenak', 'Helbide elektronikoa', 'Taldeak'])   # invented/eu
    ws.append(['Icía', 'Seoane', 'icia.seoane@example.com'])
    ws.append(['Brais', 'Gondar Veiga', 'brais.gondar@example.com'])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    d = client.post('/parse-xls',
                    data={'file': (buf, 'courseid_2_participants.xlsx')}).get_json()
    assert d['format'] == 'moodle'
    assert d['note'] == 'inferred_columns'
    assert d['students'] == ['Seoane, Icía', 'Gondar Veiga, Brais']
